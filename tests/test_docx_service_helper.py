import pytest
from docx import Document
from app.services.documents.docx_service import _set_cell_rtl

def test_set_cell_rtl_coverage():
    """تست واحد برای اطمینان از پوشش خطوط تابع _set_cell_rtl."""
    # ایجاد یک سند داکیومنت واقعی برای دریافت شیء cell
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    
    # فراخوانی تابع برای پوشش خطوط ۲۷ تا ۳۱
    try:
        _set_cell_rtl(cell)
    except Exception as e:
        pytest.fail(f"تست _set_cell_rtl با خطا مواجه شد: {e}")
    
    # بررسی اینکه آیا المان tcMar به سلول اضافه شده است
    # دسترسی به XML برای اطمینان از عملکرد صحیح
    xml_content = cell._element.xml
    assert "w:tcMar" in xml_content, "المان w:tcMar در XML سلول پیدا نشد!"
