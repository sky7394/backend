import pytest
from docx.document import Document as DocumentObject
from app.services.documents.docx_service import (
    build_exam_docx,
    generate_exam_docx_bytes,
    _set_run_font,
    DEFAULT_FONT_NAME
)
from docx import Document

@pytest.fixture
def sample_exam_data():
    return {
        "title": "آزمون فیزیک دهم",
        "grade": "دهم ریاضی",
        "subject": "فیزیک",
        "questions": [
            {
                "question_text": "یکای کار در SI کدام است؟",
                "score": 2,
                "options": ["ژول", "نیوتون", "وات", "پاسکال"],
                "correct_answer": "ژول",
                "explanation": "یکای کار و انرژی در سامانه بین‌المللی ژول (J) است."
            },
            {
                "question_text": "شتاب گرانش زمین تقریباً چند متر بر مجذور ثانیه است؟",
                "score": 1.5,
                "options": ["9.8", "10", "1", "0"],
                "correct_answer": "9.8",
                "explanation": None
            }
        ]
    }

def test_build_exam_docx_without_answers(sample_exam_data):
    doc = build_exam_docx(sample_exam_data, include_answers=False)
    assert isinstance(doc, DocumentObject)
    
    # استخراج متن تمام پاراگراف‌ها و جدول‌ها
    paragraphs_text = "\n".join([p.text for p in doc.paragraphs])
    tables_text = "\n".join([cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    full_text = tables_text + "\n" + paragraphs_text

    assert "آزمون فیزیک دهم" in full_text
    assert "یکای کار در SI کدام است؟" in full_text
    assert "[2 نمره]" in full_text
    # پاسخ‌نامه نباید باشد
    assert "کلید و پاسخ‌نامه تشریحی آزمون" not in full_text

def test_build_exam_docx_with_answers(sample_exam_data):
    doc = build_exam_docx(sample_exam_data, include_answers=True)
    paragraphs_text = "\n".join([p.text for p in doc.paragraphs])
    tables_text = "\n".join([cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    full_text = tables_text + "\n" + paragraphs_text
    
    assert "کلید و پاسخ‌نامه تشریحی آزمون" in full_text
    assert "پاسخ صحیح: ژول" in full_text
    assert "سامانه بین‌المللی ژول" in full_text

def test_generate_exam_docx_bytes(sample_exam_data):
    data_bytes = generate_exam_docx_bytes(sample_exam_data, include_answers=True)
    assert isinstance(data_bytes, bytes)
    assert len(data_bytes) > 0
    # فایل docx دارای امضای استاندارد فایل‌های zip (PK) است
    assert data_bytes.startswith(b"PK")

def test_set_run_font_xml():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("تست فونت")
    _set_run_font(r, font_name="Vazirmatn", size_pt=12, bold=True)
    
    xml = r._element.xml
    assert 'w:cs="Vazirmatn"' in xml
    assert 'w:ascii="Vazirmatn"' in xml
