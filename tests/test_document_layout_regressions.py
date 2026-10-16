from copy import deepcopy
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_DIRECTION
from PIL import Image
import pytest

from app.schemas.export import ExamExportRequestSchema
from app.services.documents.document_content import normalize_questions, resolve_header
from app.services.documents.docx_service import build_exam_docx, generate_exam_docx_bytes
from app.services.documents.exam_models import NormalizedQuestion
from app.services.documents.math_rendering import render_math, split_math
from app.services.documents import pdf_service
from app.services.documents.pdf_layout import safe_text


@pytest.fixture
def exam():
    return {
        "title": "آزمون فیزیک دهم",
        "grade": "دهم", "subject": "فیزیک",
        "header_config": {
            "school_name": "دبیرستان دانش", "grade_field": "دهم / ریاضی",
            "teacher_name": "مریم احمدی", "exam_date": "۱۴۰۵/۰۷/۲۳",
            "exam_duration": "۹۰ دقیقه", "class_name": "۱۰۱",
        },
        "questions": [
            {"question_text": r"سرعت را با رابطه $v=\frac{\Delta x}{\Delta t}$ محاسبه کنید.",
             "score": 0, "correct_answer": "ANSWER_SECRET",
             "explanation": r"از رابطه $\frac{20}{4}=5$ استفاده می‌کنیم.",
             "formula": r"E=mc^2", "answer_lines": 3},
            {"question_text": "یکای نیرو چیست؟", "options": ["ژول", "نیوتون"],
             "score": 2, "correct_answer": "نیوتون"},
        ],
    }


def test_normalization_preserves_aliases_zero_and_input():
    source = [
        {"text": "سؤال", "score": 0, "correct_answer": 0, "options": None},
        {"question": "انتخاب", "options": [{"text": "الف"}, {"text": "ب"}],
         "correct_option_index": 1, "explanation": "شرح"},
        NormalizedQuestion(question_text="سؤال", answer="پاسخ"),
    ]
    original = deepcopy(source)
    result = normalize_questions(source)
    assert result[0].points == 0 and result[0].answer == "0"
    assert result[0].options == []
    assert result[1].answer == "ب" and result[1].explanation == "شرح"
    result[2].answer = None
    assert source == original


def test_export_schema_keeps_optional_header_and_formula(exam):
    payload = ExamExportRequestSchema.model_validate(exam)
    data = payload.model_dump()
    config = resolve_header(data)
    assert config.course_name == "فیزیک"
    assert config.school_name == "دبیرستان دانش"
    assert data["questions"][0]["formula"] == "E=mc^2"
    assert data["questions"][0]["score"] == 0


def test_docx_rtl_columns_header_and_answer_sheet(exam):
    original = deepcopy(exam)
    doc = Document(BytesIO(generate_exam_docx_bytes(exam, include_answers=True)))
    header, questions, answers = doc.tables
    assert questions.table_direction == WD_TABLE_DIRECTION.RTL
    assert [c.text for c in questions.rows[0].cells] == ["ردیف", "شرح سؤال", "بارم"]
    assert questions.rows[1].cells[0].text == "1"
    assert "[0 نمره]" in questions.rows[1].cells[2].text
    assert questions.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert questions.rows[1].cells[1]._tc.xpath(".//w:pBdr/w:bottom")
    text = "\n".join(c.text for row in header.rows for c in row.cells)
    for value in ("دبیرستان دانش", "فیزیک", "دهم / ریاضی", "مریم احمدی", "۱۴۰۵/۰۷/۲۳",
                  "۹۰ دقیقه", "نام و نام خانوادگی", "شماره داوطلبی", "کلاس", "نمره به عدد",
                  "نمره به حروف", "امضای مصحح"):
        assert value in text
    assert "ANSWER_SECRET" not in questions._tbl.xml
    assert "ANSWER_SECRET" in answers._tbl.xml
    elements = list(doc.element.body)
    question_index = elements.index(questions._tbl)
    answer_index = elements.index(answers._tbl)
    assert any(e.xpath('.//w:br[@w:type="page"]') for e in elements[question_index:answer_index])
    assert len(doc.inline_shapes) == 3
    assert exam == original


def test_docx_without_answers_does_not_leak_explanations(exam):
    doc = build_exam_docx(exam, include_answers=False)
    assert len(doc.tables) == 2
    assert "ANSWER_SECRET" not in doc.element.xml
    assert "پاسخ‌نامه" not in doc.element.xml
    assert len(doc.inline_shapes) == 2


class RecordingPDF(pdf_service.ExamPDF):
    def __init__(self):
        super().__init__()
        self.cells = []

    def cell(self, w=None, h=None, text="", **kwargs):
        self.cells.append((self.page_no(), self.get_x(), self.get_y(), text))
        return super().cell(w=w, h=h, text=text, **kwargs)


def test_pdf_physical_column_order_and_answer_page(monkeypatch, exam):
    monkeypatch.setattr(pdf_service, "ExamPDF", RecordingPDF)
    original = deepcopy(exam)
    pdf = pdf_service.build_exam_pdf(exam, include_answers=True)
    labels = {name: next(c for c in pdf.cells if c[3] == safe_text(pdf, name))
              for name in ("ردیف", "شرح سؤال", "بارم")}
    assert labels["ردیف"][1] > labels["شرح سؤال"][1] > labels["بارم"][1]
    answer_pages = [c[0] for c in pdf.cells if "ANSWER_SECRET" in c[3]]
    assert answer_pages and min(answer_pages) > labels["ردیف"][0]
    assert not any("ANSWER_SECRET" in c[3] for c in pdf.cells if c[0] == 1)
    assert pdf.is_unicode_font
    assert len(pdf.image_cache.images) >= 3
    assert bytes(pdf.output()).startswith(b"%PDF")
    assert exam == original


def test_pdf_oversized_question_paginates_and_repeats_header(monkeypatch):
    monkeypatch.setattr(pdf_service, "ExamPDF", RecordingPDF)
    data = {"questions": [{"question_text": "متن بلند سؤال آزمایشی " * 650, "score": 4,
                           "correct_answer": "ANSWER_SECRET", "answer_lines": 0}]}
    pdf = pdf_service.build_exam_pdf(data, include_answers=False)
    label = safe_text(pdf, "ردیف")
    header_pages = {c[0] for c in pdf.cells if c[3] == label}
    assert pdf.page_no() >= 2
    assert header_pages == set(range(1, pdf.page_no() + 1))
    assert all(c[2] < pdf.h - pdf.b_margin for c in pdf.cells)
    assert not any("ANSWER_SECRET" in c[3] for c in pdf.cells)
    assert bytes(pdf.output()).startswith(b"%PDF")


@pytest.mark.parametrize("expression", [r"\frac{1}{2}", r"\sqrt{x^2+y^2}",
    r"F=G\frac{m_1m_2}{r^2}", r"\vec{F}=m\vec{a}", r"\alpha+\beta=\pi", "E=mc^2"])
def test_math_renders_as_one_readable_image(expression):
    result = render_math(expression)
    with Image.open(BytesIO(result.data)) as image:
        assert image.width > 5 and image.height > 5
    assert result.width_mm > 0 and result.height_mm > 0


def test_math_delimiters_and_literal_fallback(caplog):
    parts = split_math(r"متن $x^2$ و \(\frac{1}{2}\) و \[\sqrt{x}\] و $$v=t$$")
    assert [value for kind, value in parts if kind == "math"] == [
        "x^2", r"\frac{1}{2}", r"\sqrt{x}", "v=t"]
    assert ("math", "E = mc^2") in split_math("رابطه E = mc^2 را بنویسید")
    result = render_math(r"\unsupportedcommand{x}")
    assert result.data.startswith(b"\x89PNG")
    assert "literal expression" in caplog.text
